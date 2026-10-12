"""Word (.docx) <-> ProseMirror JSON, with python-docx.

Export (``pm_to_docx``): headings, paragraphs with bold / italic / underline /
strike / code / sub / sup / highlight / links, bullet, numbered and task
lists (nested), block quotes, code blocks, tables, images, callouts and
details (as boxed text), math (as monospace text). Import (``docx_to_pm``):
the same set back, from Word's own styles and numbering. Images found in a
.docx are handed to ``image_sink(blob, mime, name) -> attrs | None`` so the
caller can store them as attachments.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from core.khaidocs.html_pm import safe_href, trim_inline
from core.khaidocs.text import nfc

MONO = ("consolas", "courier", "courier new", "menlo", "monaco", "source code", "lucida console")
MAX_BODY = 200_000  # body elements read from one .docx
MAX_IMAGE_WIDTH = Inches(6)
ImageBytes = Callable[[dict], "tuple[bytes, str] | None"]
ImageSink = Callable[[bytes, str, str], "dict | None"]


# =================================================================== export


def pm_to_docx(doc: dict, *, title: str | None = None, image_bytes: ImageBytes | None = None) -> bytes:
    document = Document()
    writer = _Writer(document, image_bytes)
    if title:
        document.add_heading(nfc(title), level=0)
    writer.blocks((doc or {}).get("content") or [], 0)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class _Writer:
    def __init__(self, document, image_bytes: ImageBytes | None) -> None:
        self.document = document
        self.image_bytes = image_bytes
        self._numbering: dict[str, int] = {}

    # -- inline
    def runs(self, paragraph, nodes: list[dict] | None) -> None:
        for node in nodes or []:
            kind = node.get("type")
            if kind == "text":
                marks = {m["type"]: m for m in node.get("marks") or []}
                link = marks.get("link")
                href = safe_href((link.get("attrs") or {}).get("href")) if link else None
                target = _hyperlink(paragraph, href) if href else None
                self.run(paragraph, node.get("text", ""), marks, target)
            elif kind == "hardBreak":
                paragraph.add_run().add_break()
            elif kind == "mathInline":
                run = paragraph.add_run((node.get("attrs") or {}).get("text", ""))
                run.font.name = "Consolas"
                run.italic = True
            elif kind == "mention":
                attrs = node.get("attrs") or {}
                paragraph.add_run(("@" if attrs.get("entityType") != "page" else "") + (attrs.get("label") or ""))
            elif node.get("content"):
                self.runs(paragraph, node["content"])

    @staticmethod
    def run(paragraph, text: str, marks: dict, container=None) -> None:
        run = paragraph.add_run(text)
        if container is not None:
            container.append(run._r)
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.underline = True
        if "bold" in marks:
            run.bold = True
        if "italic" in marks:
            run.italic = True
        if "underline" in marks:
            run.underline = True
        if "strike" in marks:
            run.font.strike = True
        if "code" in marks:
            run.font.name = "Consolas"
        if "subscript" in marks:
            run.font.subscript = True
        if "superscript" in marks:
            run.font.superscript = True
        if "highlight" in marks:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    # -- blocks
    def blocks(self, nodes: list[dict], level: int, parent=None) -> None:
        for node in nodes:
            self.block(node, level, parent)

    def paragraph(self, parent, style: str | None = None):
        target = parent if parent is not None else self.document
        paragraph = target.add_paragraph()
        if style:
            try:
                paragraph.style = style
            except KeyError:
                pass
        return paragraph

    def block(self, node: dict, level: int, parent=None) -> None:
        kind = node.get("type")
        content = node.get("content") or []
        attrs = node.get("attrs") or {}
        if kind == "paragraph":
            self.runs(self.paragraph(parent), content)
        elif kind == "heading":
            heading = self.document.add_heading(level=min(max(int(attrs.get("level") or 1), 1), 9)) \
                if parent is None else self.paragraph(parent)
            self.runs(heading, content)
            if parent is not None:
                for run in heading.runs:
                    run.bold = True
        elif kind in ("bulletList", "orderedList", "taskList"):
            self.list(node, level, parent)
        elif kind == "blockquote":
            for child in content:
                if child.get("type") == "paragraph":
                    self.runs(self.paragraph(parent, "Quote"), child.get("content"))
                else:
                    self.block(child, level, parent)
        elif kind == "codeBlock":
            paragraph = self.paragraph(parent)
            text = "".join(c.get("text", "") for c in content)
            for index, line in enumerate(text.split("\n")):
                if index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            _shade_paragraph(paragraph, "F2F2F2")
        elif kind == "horizontalRule":
            _bottom_border(self.paragraph(parent))
        elif kind == "pageBreak":
            self.document.add_page_break()
        elif kind == "table":
            self.table(node, parent)
        elif kind == "image":
            self.image(node, parent)
        elif kind == "callout":
            cell = self.document.add_table(rows=1, cols=1).cell(0, 0) if parent is None else parent
            if parent is None:
                self.document.tables[-1].style = "Table Grid"
                _shade_cell(cell, {"info": "E8F1FB", "success": "E6F4EA", "warning": "FFF4D6",
                                   "danger": "FDE8E8"}.get(attrs.get("type"), "E8F1FB"))
                cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            self.blocks(content, level, cell)
            if parent is None:
                self.document.add_paragraph()
        elif kind == "details":
            summary = next((c for c in content if c.get("type") == "detailsSummary"), None)
            body = next((c for c in content if c.get("type") == "detailsContent"), None)
            paragraph = self.paragraph(parent)
            self.runs(paragraph, (summary or {}).get("content"))
            for run in paragraph.runs:
                run.bold = True
            self.blocks((body or {}).get("content") or [], level, parent)
        elif kind == "mathBlock":
            run = self.paragraph(parent).add_run(attrs.get("text", ""))
            run.font.name = "Consolas"
            run.italic = True
        elif kind in ("attachment", "video", "audio", "pdf"):
            paragraph = self.paragraph(parent)
            name = attrs.get("name") or attrs.get("src") or "file"
            self.run(paragraph, str(name), {}, None)
        elif kind in ("embed", "drawio", "excalidraw"):
            href = safe_href(attrs.get("src") or attrs.get("url") or "")
            if href:
                paragraph = self.paragraph(parent)
                self.run(paragraph, href, {"link": {}}, _hyperlink(paragraph, href))
        elif content:
            self.blocks(content, level, parent)

    def table(self, node: dict, parent) -> None:
        rows = [r.get("content") or [] for r in node.get("content") or []]
        if not rows:
            return
        width = max(sum(int((c.get("attrs") or {}).get("colspan") or 1) for c in r) for r in rows)
        host = self.document if parent is None else parent
        table = host.add_table(rows=len(rows), cols=max(width, 1))
        try:
            table.style = "Table Grid"
        except Exception:  # noqa: BLE001 - a cell has no styles part of its own
            pass
        for r_index, row in enumerate(rows):
            col = 0
            for cell_node in row:
                span = int((cell_node.get("attrs") or {}).get("colspan") or 1)
                if col >= width:
                    break
                cell = table.cell(r_index, col)
                if span > 1 and col + span - 1 < width:
                    cell = cell.merge(table.cell(r_index, col + span - 1))
                body = cell_node.get("content") or []
                first = True
                for child in body:
                    if first and child.get("type") == "paragraph":
                        self.runs(cell.paragraphs[0], child.get("content"))
                    else:
                        self.block(child, 0, cell)
                    first = False
                if cell_node.get("type") == "tableHeader":
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                col += span
        if parent is None:
            self.document.add_paragraph()

    def image(self, node: dict, parent) -> None:
        attrs = node.get("attrs") or {}
        loaded = None
        if self.image_bytes:
            try:
                loaded = self.image_bytes(attrs)
            except Exception:  # noqa: BLE001
                loaded = None
        paragraph = self.paragraph(parent)
        if loaded:
            try:
                paragraph.add_run().add_picture(io.BytesIO(loaded[0]), width=None)
                picture = paragraph.runs[-1].element.xpath(".//wp:extent")[0]
                if int(picture.get("cx")) > MAX_IMAGE_WIDTH:
                    ratio = MAX_IMAGE_WIDTH / int(picture.get("cx"))
                    for tag in (".//wp:extent", ".//a:ext"):
                        for item in paragraph.runs[-1].element.xpath(tag):
                            item.set("cx", str(int(int(item.get("cx")) * ratio)))
                            item.set("cy", str(int(int(item.get("cy")) * ratio)))
                return
            except Exception:  # noqa: BLE001 - unsupported image format
                for run in list(paragraph.runs):
                    run._r.getparent().remove(run._r)
        label = attrs.get("alt") or attrs.get("src") or "image"
        paragraph.add_run(f"[{label}]").italic = True

    def list(self, node: dict, level: int, parent) -> None:
        kind = node["type"]
        ordered = kind == "orderedList"
        num_id = self._new_num("List Number") if ordered and parent is None else None
        for item in node.get("content") or []:
            first = True
            for child in item.get("content") or []:
                if child.get("type") in ("bulletList", "orderedList", "taskList"):
                    self.list(child, level + 1, parent)
                    continue
                if child.get("type") != "paragraph":
                    self.block(child, level, parent)
                    continue
                suffix = "" if level == 0 else f" {min(level + 1, 3)}"
                style = ("List Number" if ordered else "List Bullet") + suffix
                paragraph = self.paragraph(parent, style if first else None)
                if first and kind == "taskList":
                    checked = (item.get("attrs") or {}).get("checked")
                    paragraph.add_run("☑ " if checked else "☐ ")
                self.runs(paragraph, child.get("content"))
                if first and num_id is not None:
                    _set_num(paragraph, num_id, level)
                if not first:
                    paragraph.paragraph_format.left_indent = Inches(0.25 * (level + 1))
                first = False

    def _new_num(self, style_name: str) -> int | None:
        """A fresh numbering instance for the style's list, restarting at 1."""

        try:
            numbering = self.document.part.numbering_part.element
            style = self.document.styles[style_name]
            num_pr = style.element.pPr.numPr
            base = int(num_pr.numId.val)
            base_num = next(n for n in numbering.num_lst if int(n.numId) == base)
            abstract = base_num.abstractNumId.val
            new_id = max(int(n.numId) for n in numbering.num_lst) + 1
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(new_id))
            ref = OxmlElement("w:abstractNumId")
            ref.set(qn("w:val"), str(abstract))
            num.append(ref)
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            start = OxmlElement("w:startOverride")
            start.set(qn("w:val"), "1")
            override.append(start)
            num.append(override)
            numbering.append(num)
            return new_id
        except Exception:  # noqa: BLE001 - numbering part missing: plain numbering
            return None


def _set_num(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def _hyperlink(paragraph, href: str):
    rel = paragraph.part.relate_to(
        href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), rel)
    paragraph._p.append(element)
    return element


def _shade(properties, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _shade_paragraph(paragraph, color: str) -> None:
    _shade(paragraph._p.get_or_add_pPr(), color)


def _shade_cell(cell, color: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), color)


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)


# =================================================================== import


def docx_to_pm(data: bytes, *, image_sink: ImageSink | None = None) -> dict:
    document = Document(io.BytesIO(data))
    reader = _Reader(document, image_sink)
    reader.read_body()
    return {"type": "doc", "content": reader.out or [{"type": "paragraph"}]}


class _Reader:
    def __init__(self, document, image_sink: ImageSink | None) -> None:
        self.document = document
        self.image_sink = image_sink
        self.out: list[dict] = []
        self._numfmt = self._numbering_formats()
        self.code: list[str] | None = None

    def _numbering_formats(self) -> dict[tuple[int, int], str]:
        formats: dict[tuple[int, int], str] = {}
        try:
            numbering = self.document.part.numbering_part.element
        except Exception:  # noqa: BLE001
            return formats
        abstract: dict[int, dict[int, str]] = {}
        for a in numbering.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in a.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                levels[int(lvl.get(qn("w:ilvl")))] = fmt.get(qn("w:val")) if fmt is not None else "decimal"
            abstract[int(a.get(qn("w:abstractNumId")))] = levels
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            for level, fmt in abstract.get(int(ref.get(qn("w:val"))), {}).items():
                formats[(int(num.get(qn("w:numId"))), level)] = fmt
        return formats

    # -- body
    def read_body(self) -> None:
        stack: list[tuple[int, dict]] = []  # (level, list node)
        body = self.document.element.body
        for index, child in enumerate(body.iterchildren()):
            if index > MAX_BODY:
                break
            tag = child.tag
            if tag == qn("w:p"):
                paragraph = Paragraph(child, self.document._body)
                self.paragraph(paragraph, stack)
            elif tag == qn("w:tbl"):
                self.close_code()
                stack.clear()
                table = self.table(Table(child, self.document._body))
                if table:
                    self.out.append(table)
        self.close_code()

    def close_code(self) -> None:
        if self.code is not None:
            text = "\n".join(self.code)
            node: dict = {"type": "codeBlock", "attrs": {"language": None}}
            if text:
                node["content"] = [{"type": "text", "text": text}]
            self.out.append(node)
            self.code = None

    def style_name(self, paragraph) -> str:
        try:
            return (paragraph.style.name or "") if paragraph.style is not None else ""
        except Exception:  # noqa: BLE001
            return ""

    def list_info(self, paragraph) -> tuple[bool, int, bool] | None:
        """(is a list item, level, ordered)."""

        p_pr = paragraph._p.pPr
        num_pr = p_pr.numPr if p_pr is not None else None
        style = self.style_name(paragraph)
        if num_pr is None and paragraph.style is not None:
            style_ppr = paragraph.style.element.pPr
            num_pr = style_ppr.numPr if style_ppr is not None else None
        explicit = p_pr is not None and p_pr.numPr is not None and p_pr.numPr.ilvl is not None
        if num_pr is not None and num_pr.numId is not None and int(num_pr.numId.val) != 0:
            level = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
            if not explicit and re.match(r"list (bullet|number)", style.lower()):
                level = _style_level(style)
            num_id = int(num_pr.numId.val)
            fmt = self._numfmt.get((num_id, level))
            if fmt is None and style.lower().startswith("list"):
                fmt = "bullet" if "bullet" in style.lower() else "decimal"
            return True, level, fmt not in (None, "bullet", "none")
        if style.lower().startswith("list bullet"):
            return True, _style_level(style), False
        if style.lower().startswith("list number"):
            return True, _style_level(style), True
        return None

    def paragraph(self, paragraph: Paragraph, stack: list) -> None:
        style = self.style_name(paragraph)
        inline = self.inline(paragraph)
        text_only = "".join(
            n.get("text", "") if n["type"] == "text" else "\n" for n in inline if n["type"] in ("text", "hardBreak")
        )
        # code: every text run monospace, or a code style
        if inline and all(
            n["type"] == "hardBreak"
            or (n["type"] == "text" and any(m["type"] == "code" for m in n.get("marks", [])))
            for n in inline
        ) or style.lower() in ("code", "source code", "html preformatted"):
            stack.clear()
            if self.code is None:
                self.code = []
            self.code.append(text_only)
            return
        self.close_code()
        lowered = style.lower()
        info = self.list_info(paragraph)
        if not inline and not any(n["type"] == "image" for n in inline):
            return
        images = [n for n in inline if n["type"] == "image"]
        inline = [n for n in inline if n["type"] != "image"]
        inline = trim_inline(inline)
        if info and inline:
            _, level, ordered = info
            kind, checked = ("orderedList" if ordered else "bulletList"), None
            first = inline[0]
            if first["type"] == "text" and first["text"][:1] in ("\u2610", "\u2611") and not ordered:
                kind, checked = "taskList", first["text"][0] == "\u2611"
                rest = first["text"][1:].lstrip()
                inline = ([{**first, "text": rest}] if rest else []) + inline[1:]
            self.add_list_item(stack, level, kind, inline, checked)
            self.out_extend_images(images, stack)
            return
        stack.clear()
        heading = re.match(r"heading (\d)", lowered)
        if heading or lowered == "title":
            node: dict = {"type": "heading", "attrs": {"level": min(int(heading.group(1)), 6) if heading else 1}}
            if inline:
                node["content"] = inline
            self.out.append(node)
        elif lowered in ("quote", "intense quote", "block text"):
            if self.out and self.out[-1]["type"] == "blockquote":
                self.out[-1]["content"].append({"type": "paragraph", "content": inline})
            else:
                self.out.append({"type": "blockquote", "content": [{"type": "paragraph", "content": inline}]})
        elif inline:
            self.out.append({"type": "paragraph", "content": inline})
        self.out += images

    def out_extend_images(self, images: list[dict], stack: list) -> None:
        if images:
            stack.clear()
            self.out += images

    def add_list_item(self, stack: list, level: int, kind: str, inline: list[dict],
                      checked: bool | None = None) -> None:
        body = [{"type": "paragraph", "content": inline}] if inline else [{"type": "paragraph"}]
        item = ({"type": "taskItem", "attrs": {"checked": bool(checked)}, "content": body}
                if kind == "taskList" else {"type": "listItem", "content": body})
        while stack and stack[-1][0] > level:
            stack.pop()
        if stack and stack[-1][0] == level and stack[-1][1]["type"] != kind:
            stack.pop()
        if stack and stack[-1][0] == level:
            stack[-1][1]["content"].append(item)
            return
        node: dict = {"type": kind, "content": [item]}
        if kind == "orderedList":
            node["attrs"] = {"start": 1}
        if stack:
            stack[-1][1]["content"][-1]["content"].append(node)
        else:
            self.out.append(node)
        stack.append((level, node))

    # -- inline
    def inline(self, paragraph: Paragraph) -> list[dict]:
        out: list[dict] = []
        for child in paragraph._p.iterchildren():
            if child.tag == qn("w:r"):
                out += self.run(Run(child, paragraph), None)
            elif child.tag == qn("w:hyperlink"):
                href = None
                rid = child.get(qn("r:id"))
                if rid and rid in paragraph.part.rels:
                    rel = paragraph.part.rels[rid]
                    href = safe_href(rel.target_ref) if rel.is_external else None
                anchor = child.get(qn("w:anchor"))
                if anchor and not href:
                    href = None
                for run in child.findall(qn("w:r")):
                    out += self.run(Run(run, paragraph), href)
        return _merge(out)

    def run(self, run: Run, href: str | None) -> list[dict]:
        marks: list[dict] = []
        try:
            if run.bold:
                marks.append({"type": "bold"})
            if run.italic:
                marks.append({"type": "italic"})
            if run.underline and not href:
                marks.append({"type": "underline"})
            if run.font.strike:
                marks.append({"type": "strike"})
            if run.font.subscript:
                marks.append({"type": "subscript"})
            if run.font.superscript:
                marks.append({"type": "superscript"})
            if run.font.highlight_color:
                marks.append({"type": "highlight"})
            name = (run.font.name or "").lower()
        except Exception:  # noqa: BLE001
            name = ""
        if name in MONO:
            marks = [{"type": "code"}]
        if href:
            marks.append({"type": "link", "attrs": {"href": href, "internal": False}})
        out: list[dict] = []
        buffer: list[str] = []

        def flush() -> None:
            if buffer:
                text = nfc("".join(buffer))
                buffer.clear()
                if text:
                    node: dict = {"type": "text", "text": text}
                    if marks:
                        node["marks"] = [dict(m) for m in marks]
                    out.append(node)

        for child in run._r.iterchildren():
            tag = child.tag
            if tag == qn("w:t"):
                buffer.append(child.text or "")
            elif tag == qn("w:tab"):
                buffer.append("\t")
            elif tag == qn("w:br") or tag == qn("w:cr"):
                if child.get(qn("w:type")) in (None, "textWrapping"):
                    flush()
                    out.append({"type": "hardBreak"})
            elif tag == qn("w:drawing") or tag == qn("w:pict"):
                flush()
                image = self.image(child, run)
                if image:
                    out.append(image)
        flush()
        return out

    def image(self, element, run: Run) -> dict | None:
        if self.image_sink is None:
            return None
        blips = element.xpath(".//a:blip")
        if not blips:
            return None
        rid = blips[0].get(qn("r:embed"))
        part = run.part.related_parts.get(rid) if rid else None
        if part is None:
            return None
        alt = ""
        props = element.xpath(".//wp:docPr")
        if props:
            alt = props[0].get("descr") or props[0].get("title") or ""
        name = str(part.partname).rsplit("/", 1)[-1]
        attrs = self.image_sink(part.blob, part.content_type, name)
        if not attrs:
            return None
        if alt:
            attrs = {**attrs, "alt": nfc(alt)}
        return {"type": "image", "attrs": attrs}

    # -- tables
    def table(self, table: Table) -> dict | None:
        rows = []
        for row in table.rows:
            cells = []
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                tc_pr = cell._tc.tcPr
                span = tc_pr.gridSpan if tc_pr is not None and tc_pr.gridSpan is not None else 1
                inner = _Reader(self.document, self.image_sink)
                inner._numfmt = self._numfmt
                stack: list = []
                for p in cell._tc.iterchildren(qn("w:p")):
                    inner.paragraph(Paragraph(p, cell), stack)
                inner.close_code()
                node: dict = {"type": "tableCell", "content": inner.out or [{"type": "paragraph"}]}
                if span and span > 1:
                    node["attrs"] = {"colspan": span}
                cells.append(node)
            if cells:
                rows.append({"type": "tableRow", "content": cells})
        if not rows:
            return None
        for cell in rows[0]["content"]:
            cell["type"] = "tableHeader"
        return {"type": "table", "content": rows}


def _style_level(style: str) -> int:
    match = re.search(r"(\d+)$", style)
    return max(int(match.group(1)) - 1, 0) if match else 0


def _merge(nodes: list[dict]) -> list[dict]:
    out: list[dict] = []
    for node in nodes:
        if out and node["type"] == "text" and out[-1]["type"] == "text" \
                and out[-1].get("marks") == node.get("marks"):
            out[-1] = {**out[-1], "text": out[-1]["text"] + node["text"]}
        else:
            out.append(node)
    return out
